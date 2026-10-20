"""Create the editable project proposal using the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'docx'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.78)
sec.footer_distance = Inches(.32)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Subtitle'].font.size = Pt(13)
for name, size in [('Heading 1', 19), ('Heading 2', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('Show Me the Leak  |  Proposal draft  |  ')
run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Show Me the Leak Project Proposal'
doc.core_properties.subject = 'Multimodal image privacy detection and redaction'
doc.core_properties.author = ''
doc.core_properties.keywords = 'PII, visual privacy, VLM, OCR, redaction, benchmark'

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=2):
    return doc.add_heading(text, level)

def page(title):
    doc.add_page_break()
    h(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]: el.set(qn('w:' + key), val)
        borders.append(el)
    props.append(borders)
    for i, values in enumerate([headers] + rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), '203C50' if i == 0 else ('F1F5F7' if i % 2 == 0 else 'FFFFFF')); pr.append(shading)
            mar = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '90'); el.set(qn('w:type'), 'dxa'); mar.append(el)
            pr.append(mar)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.line_spacing = 1.03
            if widths[j] < 1.05: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(text)); run.font.size = Pt(10)
            run.bold = i == 0
            run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '000000')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('Show Me the Leak', 'Title')
doc.add_paragraph('Project proposal for multimodal image privacy detection and redaction', 'Subtitle')
p('Draft 0.1  |  9 October 2026')
p('Prepared by  [Student or team name]\nInstitution  [Institution or department]\nSupervisor  [Supervisor name]')
h('Proposal summary')
p('We propose to develop and evaluate a system that identifies personally identifiable information (PII) and other privacy-sensitive content in images, locates the relevant regions, and creates a redacted copy for human review. The project combines vision-language models with optical character recognition (OCR) to examine documents and everyday photographs. Its central question is whether these models can detect and localize enough sensitive content to support a useful privacy review workflow.')
p('The project already has a notebook prototype and a 50-image pilot benchmark. Six model families are included in the implementation, although two did not complete inference in the recorded run. The next stage will repair and verify those runs, improve the evaluation data, select a suitable model or combination, and build a demonstrable application. Training will be considered after the baseline study identifies specific weaknesses.')
h('Why this project is needed')
p('An image can reveal an address on an envelope, an identifier on a card, a vehicle registration plate, a face, or private text on a screen. A person may notice the main subject while overlooking these details. Manual review is difficult to repeat consistently, especially when images contain many small regions or dense text.')
p('Text recognition can recover written content, but it does not by itself decide which content is private. General visual models can recognize broader context, yet may miss regions, produce incorrect boxes, or fail to follow an output format. A useful system must therefore measure detection, localization, redaction quality, and operational failures separately. The VISPR research provides the starting taxonomy and image-level privacy task [1].')
h('Expected outcome')
p('The expected deliverable is a research prototype with reproducible model comparisons, a reviewed annotation subset, an editable redaction interface, and a documented deployment package. It will support privacy review; it will not certify that an image is anonymous or free of all sensitive information.')

page('1 Objectives and project scope')
h('Research objectives')
p('The first objective is to compare six pretrained models on the same images and privacy taxonomy. The comparison will measure which privacy attributes they detect, which they miss, how reliably they return usable outputs, and whether performance changes between documents and general images.')
p('The second objective is to evaluate where the sensitive content appears and whether the proposed redactions cover it. The third is to determine whether targeted adaptation or a combination of OCR and visual detection improves the baseline enough to justify its additional cost. The final objective is to deploy the selected approach as a reviewable prototype.')
h('Planned user capabilities')
table(['Capability', 'Expected behavior'], [
    ['Image inspection', 'Accept an image and identify supported privacy categories with visible evidence.'],
    ['Region review', 'Display bounding boxes and allow the user to add, remove, or adjust regions.'],
    ['Redaction and export', 'Apply solid masks and export a new image together with a structured report.'],
    ['Leakage review', 'Flag potential sensitive content that remains after redaction for further review.'],
    ['Benchmark reporting', 'Compare models, show missed labels for each image, and summarize failures and runtime.'],
], [1.55, 5.16])
h('Data and taxonomy')
p('The current data source is VISPR train2017 with per-image JSON annotations and the project taxonomy crosswalk. The taxonomy contains 68 entries: 67 privacy attributes and one safe label. These include direct identifiers, document categories, biometric information, and contextual or sensitive attributes. The safe label will not be treated as a positive PII detection.')
p('The current pilot contains 50 reproducibly selected images: 25 document images and 25 general images. Document membership is derived from the dataset\'s document labels, so it is a benchmark grouping rather than an independently verified document classifier. Both groups will be used for every model, including the OCR models.')
h('Boundaries')
p('The initial application will process still images. Video, live surveillance, identity recognition, and comprehensive PDF workflows are outside the first release. Sensitive traits such as religion or ethnicity will not be inferred from a person\'s appearance. Where relevant, the system will rely on explicit visible evidence. Dataset labels that conflict with this policy will be reported as a separate evaluation scope rather than silently changing the task.')

page('2 Current progress and preliminary evidence')
p('The prototype includes dataset discovery, reproducible sampling, model adapters, output validation, coordinate conversion, box previews, redaction, and per-image comparison against the annotation JSON. Recorded outputs have been reviewed without changing the original runs. No model has yet been fine-tuned for this project, and no production service has been deployed [8].')
table(['Model', 'Valid images', 'Recorded pilot status'], [
    ['Qwen3 VL 8B', '50 / 50', 'Completed after configured retries.'],
    ['Qwen2.5 VL 7B', '45 / 50', 'Five outputs failed schema or evidence validation.'],
    ['Llama 3.2 Vision 11B', '5 / 50', 'Most outputs were descriptions rather than valid structured predictions.'],
    ['Chandra 9B', '46 / 50', 'Four failures associated with incomplete generation at the token limit.'],
    ['dots.ocr 3B', '0 / 50', 'Stopped during environment setup.'],
    ['Florence 2 large ft', '0 / 50', 'Stopped during dependency or import checks.'],
], [1.58, .93, 4.20])
h('What the pilot shows')
p('On the same 45 successfully processed images, Qwen3 achieved a label-level micro F1 of 0.157 and recall of 0.088, compared with 0.060 and 0.031 for Qwen2.5. Qwen3 missed 206 image-attribute occurrences; Qwen2.5 missed 219. These counts describe missing labels across images, not the number of individual private objects [8].')
p('The low recall shows that this prototype needs further work before unattended use. The full taxonomy also contains contextual attributes that OCR rules cannot detect and sensitive attributes the current policy intentionally avoids inferring. Results must therefore be separated by task scope. The present evidence does not support a complete six-model ranking.')
h('Immediate work remaining')
p('The notebook has been revised to use isolated model environments and clearer import diagnostics. Those changes still need verification on a GPU. The next run will retry dots.ocr and Florence, investigate Llama\'s structured-output failures, and evaluate any revised Chandra generation limit as a separately versioned experiment.')
h('Limits of the present evidence')
p('The 50 images come from the training folder and form an exploratory pilot, not a held-out generalization test. Existing JSON annotations provide image-level labels, not reference boxes or text transcripts. Consequently, the current results cannot establish bounding-box IoU, OCR transcription accuracy, or the completeness of redaction. A second pass by the same model may repeat its original mistakes.')

page('3 Model selection and system method')
p('The study will keep the requested checkpoints identifiable and compare their actual roles. The general vision-language models will predict privacy labels and locations. OCR models will recognize text and layout; a shared mapping stage will assign privacy labels. OCR results will therefore be described as OCR-plus-mapping pipelines [2-7].')
table(['Candidate', 'Role and output handling'], [
    ['Qwen3 VL 8B Instruct', 'Primary visual baseline; validate requested labels and boxes from structured output.'],
    ['Qwen2.5 VL 7B Instruct', 'Comparison with the earlier Qwen vision-language model using the same task.'],
    ['Llama 3.2 Vision 11B Instruct', 'Alternative visual model; explicitly measure format compliance and failure rate.'],
    ['Chandra 9B', 'Parse recognized content and region information from its native document output.'],
    ['dots.ocr 3B', 'Parse native layout output and transform boxes from the processed image coordinates.'],
    ['Florence 2 large ft', 'Use OCR_WITH_REGION and its native postprocessor to obtain text and region geometry.'],
], [2.0, 4.71])
h('Processing sequence')
p('The system will validate the input, normalize orientation, record the image dimensions, and invoke the selected model. A model-specific adapter will convert its output into a common record containing the image identifier, attribute identifier, region coordinates, source model, and validation status. Missing or invalid boxes will remain explicit rather than being replaced with a guessed region.')
p('For Florence, the system will use its supported task token and output parser instead of assuming it accepts or produces the project\'s JSON schema. The other OCR adapters will likewise preserve native output before conversion. All boxes will be mapped back to the original image dimensions, with checks for rotation, resizing, and out-of-bounds coordinates.')
p('The redaction stage will draw opaque masks over approved regions, export a new image, and strip metadata from the export. A review step will display the result alongside the detection report. During experiments, annotation JSON will be used only by the evaluator after inference; it will not be supplied as the answer to the model.')
h('Selection decision')
p('Selection will prioritize missed sensitive content, localization quality, valid-output rate, and redaction utility. Latency, peak GPU memory, deployment complexity, and model terms will also matter. A combined OCR and visual approach will be retained only if measured gains justify its additional runtime and errors. Model size or general leaderboard claims alone will not determine the choice.')

page('4 Dataset preparation and training plan')
h('Separate development from final testing')
p('The first milestone will preserve the existing 50-image sample for reproducibility. A larger study will then use separate training, validation, and test partitions, with exact and perceptual duplicate checks before assignment. Any officially designated evaluation split will be preferred where available and suitable. Otherwise, a documented internal split will be created without claiming an official benchmark result.')
p('The proposed planning target is at least 200 held-out evaluation images, approximately balanced between document and general images where the data permits. This expands the pilot rather than replacing its records. Rare labels will be tracked explicitly; a 200-image study will not necessarily provide reliable estimates for every one of the 67 privacy attributes. Sample size will be reviewed against annotation effort and GPU availability before the final test is frozen.')
h('Add the references that localization requires')
p('A proposed subset of at least 100 held-out images will receive manually verified boxes for every in-scope visible privacy instance. Text-bearing regions in this subset will also receive transcription references for OCR evaluation. Two reviewers will annotate or independently check the subset, reconcile disagreements, and document ambiguous or intentionally excluded attributes. Training boxes and examples will come from separate images.')
p('Annotation guidance will define the tightest practical region, treatment of repeated identifiers, partial visibility, document-level versus field-level labels, and overlapping categories. Complete annotation within the selected scope is necessary to interpret unmatched predictions. The project will retain label-level results for all evaluable images and location-level results only for the verified subset.')
h('Adapt only after a stable baseline')
p('Training will begin with error analysis rather than training all six models. The team will first improve prompts, parsing, and the shared text-to-privacy mapping using development data. If important errors remain, one or two suitable candidates will undergo parameter-efficient adaptation, such as LoRA or QLoRA where supported by the architecture, software stack, and model terms.')
p('Training examples will teach supported attribute labels, valid region coordinates, and compliant output structure. Synthetic identifiers may augment scarce text examples, but their realism and label correctness must be reviewed. Class balancing and difficult negative examples will address rare attributes and unnecessary redactions. Raw sensitive values will be minimized in stored training artifacts.')
h('Controlled comparison')
p('The study will compare the original model, revised prompting, mapping changes, optional fine-tuning, and any hybrid pipeline separately. Validation data will guide checkpoint selection and thresholds. The frozen test set will be used after those decisions. A training branch will be abandoned if it does not improve the agreed privacy and utility measures enough to justify its resource cost.')

page('5 Evaluation and testing strategy')
p('Every run will record the checkpoint revision, prompt, dependency profile, image manifest, resolution policy, decoding settings, quantization, retries, and hardware. Changed settings will produce a new protocol version. Reports will include both model quality and the fraction of attempted images that produced a valid result.')
table(['Evaluation area', 'Measures and required references'], [
    ['PII label detection', 'Micro precision, recall and F1; supported macro F1; per-attribute recall; missed and extra labels. Uses image-level annotation JSON.'],
    ['Localization', 'Intersection over union (IoU) measures box overlap. Use same-label one-to-one matching at IoU at least 0.5, box precision, recall and F1. Requires verified boxes.'],
    ['OCR text quality', 'Character error rate, word error rate and optional ROUGE-L against independent transcripts. Uses a common text-normalization policy.'],
    ['Redaction quality', 'Sensitive-region pixel coverage, remaining exposed instances, unnecessary masked area, and human review of readability. Requires complete in-scope annotations.'],
    ['Operational reliability', 'Valid-output rate, load failures, parsing failures, truncation, median and tail latency, and peak GPU memory.'],
], [1.48, 5.23])
h('Make comparisons fair')
p('Primary reports will separate document and general images and show both the full privacy taxonomy and the shared text-label scope. Head-to-head quality comparisons will use the same successfully processed images and report that sample size. A separate all-attempts analysis will count failures as unavailable protection so that a model cannot appear strong merely by failing on difficult inputs. Missing annotations will be distinguished from confirmed false detections.')
p('Where sample size permits, image-level paired bootstrap intervals will describe uncertainty in differences. Results on rare labels will show their support. Sensitive-label policy exclusions will be visible. OCR transcription metrics will not be calculated on unordered PII label JSON, and high text similarity will not be treated as evidence of effective redaction.')
h('Test the complete workflow')
p('Software checks will cover each parser, coordinate conversion, mask placement, reproducible sampling, resume behavior, and missing outputs. End-to-end tests will include rotated images, small or blurred text, repeated identifiers, crowded scenes, no-PII images, invalid files, and memory failures. Failures must display a review-required state rather than silently returning an image marked safe.')
p('Redacted images will receive an independent review on the reference subset, using human checks and, where useful, a separate OCR or detection pass. The existing same-model second pass will remain a diagnostic aid. Acceptance targets for recall, exposed regions, over-redaction, and response time will be fixed after the repaired baseline and before final testing, with any unresolved misses documented.')

page('6 Deployment and planned application')
h('Prototype architecture')
p('The proposed application will have a browser interface, an inference API, and a GPU worker. A practical implementation may use FastAPI for requests and a lightweight web interface for image review. The selected framework will be confirmed after the notebook pipeline is stable. Model-specific dependency environments will remain isolated, and only the selected serving model or pipeline will need to be loaded for ordinary use.')
p('The interface will allow a user to upload an image, choose supported redaction categories, inspect proposed regions, make corrections, and export the redacted copy. Reports will show category and location without unnecessarily repeating the private values. A result with incomplete processing or unlocalized detections will clearly request further review.')
h('Deployment stages')
table(['Stage', 'Deliverable and verification'], [
    ['Local demonstration', 'GPU-backed notebook or local service with a repeatable setup and a small reviewed example set.'],
    ['Packaged prototype', 'Versioned container or environment package, model configuration, API and interface, with a repeatable installation guide.'],
    ['Controlled pilot', 'Restricted users, representative images, monitored failures and latency, and a documented rollback path.'],
], [1.46, 5.25])
h('Privacy controls')
p('Sensitive images should be processed locally or on an approved private GPU service where practical. The prototype will define file retention, deletion, access controls, and whether raw model outputs may be stored. Ordinary service logs should record timing and errors without image contents or recognized private text. Any research archive containing source images or OCR values will have separate, restricted access.')
p('Export checks will verify that masks replace the intended pixels, that the source image is not embedded in the exported result, and that metadata is removed as configured. Corrections used for future training will require a deliberate retention decision. Model and dataset terms will be reviewed before distributing checkpoints or offering access beyond the research setting.')
h('Monitoring and handover')
p('The pilot will monitor failed requests, processing time, corrections made by reviewers, and categories with repeated misses. A change to a model, prompt, parser, or redaction policy will trigger regression testing against a fixed reference set. The handover will include known limitations, supported input types, resource requirements, and instructions for stopping or reverting a faulty release.')
p('Public hosting and autonomous bulk redaction are later decisions. The initial deployment goal is a working demonstration with human review and reproducible evidence of its limits.')

page('7 Milestones resources and risks')
p('The schedule below proposes 12 weeks from the start of the next development phase. It assumes access to a compatible GPU and reviewers for annotations. It is a planning estimate; runtime repairs and annotation effort are the main dependencies. The existing pilot is completed groundwork rather than a future milestone.')
table(['Weeks', 'Work package', 'Completion evidence'], [
    ['1 to 2', 'Repair runtimes and repeat baselines', 'All six candidates attempted on the fixed sample; successful results or reproducible failure records.'],
    ['3 to 4', 'Prepare evaluation data and references', 'Frozen partitions, duplicate checks, annotation guide and reviewed box and transcript subset.'],
    ['5 to 6', 'Compare models and select candidates', 'Paired results, document/general breakdowns, error analysis and selection rationale.'],
    ['7 to 8', 'Targeted adaptation and ablations', 'Prompt, mapping and optional training experiments evaluated on validation data.'],
    ['9 to 10', 'Build and package the application', 'Working upload, review, edit and export flow; installation and privacy controls verified.'],
    ['11 to 12', 'Final testing and project report', 'Held-out evaluation, controlled demonstration, limitations, documentation and handover.'],
], [.70, 2.24, 3.77])
h('Resources and responsibilities')
p('The project needs GPU access suitable for the selected checkpoints, storage for model weights and restricted artifacts, and time for annotation and review. The previous pilot used an NVIDIA L4; larger or repeated training jobs may need additional memory or runtime. A hardware smoke test will determine feasible batch sizes before longer runs. Costs will be estimated from measured runtime and the chosen provider rather than an assumed fixed budget.')
p('The development role will own adapters, experiments and deployment; an evaluation role will manage partitions and references; a reviewer or supervisor will resolve annotation disputes and assess the final evidence. One person may perform multiple roles, but independent review of the reference subset should be retained.')
h('Main risks and responses')
p('Dependency conflicts will be addressed through isolated, versioned environments and import checks. Missing boxes or transcripts will be addressed through scoped manual annotation. Weak recall will trigger category-level error analysis and human review; excess masking will be measured as a utility cost. Dataset imbalance and policy differences will be reported through support counts and separate evaluation scopes. Hardware limits may reduce the number of training experiments, but should not remove recorded failures from the comparison.')

page('8 Deliverables and references')
h('Deliverables and completion criteria')
p('The project will deliver a reproducible benchmark notebook and model adapters; versioned manifests and an annotation guide; a reviewed reference subset; model comparison and error-analysis reports; an optional adapted checkpoint where justified; and a packaged image-review application with setup and operation instructions.')
p('Completion requires reproducible experiments, transparent reporting of failed models, evaluation on a frozen test partition, verified geometry and redaction behavior, and a working reviewed export flow. The final report will explain the chosen model or pipeline, quantify its limitations, and state whether adaptation improved the baseline. Unmet performance targets will remain visible rather than being replaced by claims of complete privacy protection.')
h('Research and model references')
p('The sources below identify the dataset background and requested model families. Model cards describe general capabilities and usage; the project\'s own experiments determine performance on its privacy task. Web sources were checked on 9 October 2026.')
refs = [
    ('1', 'Orekondy, Schiele and Fritz. Towards a Visual Privacy Advisor Understanding and Predicting Privacy Risks in Images. ICCV 2017.', 'https://arxiv.org/abs/1703.10660'),
    ('2', 'Qwen3 visual baseline. Qwen/Qwen3-VL-8B-Instruct.', 'https://huggingface.co/Qwen/Qwen3-VL-8B-Instruct'),
    ('3', 'Qwen2.5 visual comparison. Qwen/Qwen2.5-VL-7B-Instruct.', 'https://huggingface.co/Qwen/Qwen2.5-VL-7B-Instruct'),
    ('4', 'Meta visual comparison. meta-llama/Llama-3.2-11B-Vision-Instruct.', 'https://huggingface.co/meta-llama/Llama-3.2-11B-Vision-Instruct'),
    ('5', 'Original Chandra 9B OCR checkpoint. datalab-to/chandra.', 'https://huggingface.co/datalab-to/chandra'),
    ('6', 'Original dots.ocr 3B checkpoint. Project identifier rednote-hilab/dots.ocr; the model page currently redirects to dots-studio/dots.ocr.', 'https://huggingface.co/rednote-hilab/dots.ocr'),
    ('7', 'Florence OCR and region extraction. microsoft/Florence-2-large-ft.', 'https://huggingface.co/microsoft/Florence-2-large-ft'),
]
for num, text, url in refs:
    para = p(f'[{num}] {text}\n{url}')
    para.paragraph_format.space_after = Pt(6)
    for run in para.runs: run.font.size = Pt(9)
para = p('[8] Show Me the Leak project records. Saved 50-image pilot in Leak_Benchmark_Results; review in outputs/benchmark_review/README.md and associated provenance and metrics files. Sample fingerprint prefix 1df1d1c206ae. Internal project evidence, October 2026.')
for run in para.runs: run.font.size = Pt(9)

target = OUT / 'Show_Me_the_Leak_Project_Proposal_Draft.docx'
doc.save(target)
print(target)
print('Words:', sum(len(x.text.split()) for x in doc.paragraphs) + sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells))
